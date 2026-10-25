#!/usr/bin/env python3
"""Integration check for the engine shared by MergeGuru Lite and Pro."""

import shutil
import sys
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import mergeguru_engine as engine  # noqa: E402


def main():
    work = Path(tempfile.mkdtemp(prefix="mergeguru_engine_"))
    try:
        reference = work / "reference.csv"
        shuffled = work / "shuffled.csv"
        pd.DataFrame(
            [
                {"id": 1, "region": "South", "amount": 100},
                {"id": 2, "region": "North", "amount": 200},
            ]
        ).to_csv(reference, index=False)
        pd.DataFrame(
            [
                {"amount": 300, "region": "South", "id": 3, "note": "extra"},
                {"amount": 400, "region": "West", "id": 4, "note": "extra"},
            ]
        )[["amount", "region", "id", "note"]].to_csv(shuffled, index=False)

        sources, errors = engine.discover_sources([str(reference), str(shuffled)])
        assert not errors
        assert len(sources) == 2

        validation = engine.validate_sources(
            sources, 0, allow_extra=True
        )
        assert len(validation["passed_sources"]) == 2
        assert validation["extra_columns"] == ["note"]

        combined, extras, _ = engine.combine_sources(
            validation["passed_sources"],
            sources[0],
            validation["reference_columns"],
            allow_extra=True,
            extra_columns=validation["extra_columns"],
            fast_mode=True,
            optimize_memory=True,
        )
        assert len(combined) == 4
        assert list(combined.columns) == [
            "id", "region", "amount", "note", "source_file"
        ]
        assert extras == ["note"]

        filtered = engine.apply_filters(
            combined,
            [{
                "column": "region",
                "mode": "text",
                "operator": "equals",
                "value": "South",
            }],
        )
        assert len(filtered) == 2

        output = work / "result.csv"
        engine.save_dataframe(filtered, str(output), "CSV")
        assert output.exists()
        assert len(pd.read_csv(output)) == 2

        protected_output = work / "protected.csv"
        engine.save_dataframe(
            filtered,
            str(protected_output),
            "CSV-Unique text",
            unique_columns=["id", "region"],
        )
        protected = pd.read_csv(protected_output, dtype=str)
        assert protected["id"].str.startswith("'").all()
        assert protected["region"].str.startswith("'").all()
        assert not protected["amount"].str.startswith("'").any()

        excel_output = work / "streamed.xlsx"
        excel_progress = []
        excel_frame = pd.DataFrame(
            {
                "id": ["001", "002"],
                "created": [pd.Timestamp("2026-10-11"), pd.NaT],
                "literal": ["=1+1", "https://example.invalid/not-a-link"],
            }
        )
        engine.save_dataframe(
            excel_frame,
            str(excel_output),
            "Excel",
            validation_report_df=pd.DataFrame({"Status": ["PASS"]}),
            progress_callback=lambda percent, _message: excel_progress.append(percent),
        )
        workbook = load_workbook(excel_output, read_only=True, data_only=False)
        try:
            assert workbook.sheetnames == ["Data", "Validation Report"]
            rows = list(workbook["Data"].iter_rows(values_only=False))
            assert rows[1][0].value == "001"
            assert rows[1][2].value == "=1+1"
            assert rows[1][2].data_type == "s"
            assert rows[2][1].value is None
        finally:
            workbook.close()
        assert engine.EXCEL_EXPORT_ENGINE == "xlsxwriter-streaming"
        assert excel_progress[-1] == 100

        xlsx_reference = work / "reference_dates.xlsx"
        xlsx_daily = work / "daily_dates.xlsx"
        with pd.ExcelWriter(xlsx_reference, engine="openpyxl") as writer:
            pd.DataFrame(
                [{"id": 10, "region": "East", "amount": 500}]
            ).to_excel(writer, sheet_name="Dump 01-Jul-2026", index=False)
        with pd.ExcelWriter(xlsx_daily, engine="openpyxl") as writer:
            pd.DataFrame({"ignore": [1]}).to_excel(
                writer, sheet_name="Summary", index=False
            )
            pd.DataFrame(
                [{"amount": 600, "region": "East", "id": 11}]
            )[["amount", "region", "id"]].to_excel(
                writer, sheet_name="Dump 02-Jul-2026", index=False
            )
        xlsx_sources, errors = engine.discover_sources(
            [str(xlsx_reference), str(xlsx_daily)]
        )
        assert not errors
        xlsx_validation = engine.validate_sources(
            xlsx_sources,
            0,
            selected_sheet="Dump 01-Jul-2026",
            auto_match_sheet=True,
        )
        assert len(xlsx_validation["passed_sources"]) == 2
        assert xlsx_validation["source_sheet_map"][
            engine.source_cache_key(xlsx_sources[1])
        ] == "Dump 02-Jul-2026"
        xlsx_combined, _, _ = engine.combine_sources(
            xlsx_validation["passed_sources"],
            xlsx_sources[0],
            xlsx_validation["reference_columns"],
            selected_sheet="Dump 01-Jul-2026",
            source_sheet_map=xlsx_validation["source_sheet_map"],
        )
        assert len(xlsx_combined) == 2
        assert set(xlsx_combined["id"]) == {10, 11}
        print(f"shared engine check ok: {work}")
    finally:
        if "--keep" not in sys.argv:
            shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    main()
