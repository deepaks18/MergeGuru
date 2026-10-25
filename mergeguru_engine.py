#!/usr/bin/env python3
"""Shared processing engine for MergeGuru Lite and MergeGuru Pro."""

import math
import os
import tempfile
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from io import BytesIO

import pandas as pd
import psutil
import xlsxwriter
from openpyxl import load_workbook


APP_VERSION = "v3.1"
EXCEL_HARD_ROW_LIMIT = 1_048_576
EXCEL_HARD_COLUMN_LIMIT = 16_384
EXCEL_EXPORT_ENGINE = "xlsxwriter-streaming"
_SHEET_NAME_CACHE = {}

EXPORT_FORMATS = {
    "Excel": {
        "extension": ".xlsx",
        "filetypes": [("Excel files", "*.xlsx"), ("All files", "*.*")],
    },
    "CSV": {
        "extension": ".csv",
        "filetypes": [("CSV files", "*.csv"), ("All files", "*.*")],
    },
    "CSV-Unique text": {
        "extension": ".csv",
        "filetypes": [("CSV files", "*.csv"), ("All files", "*.*")],
    },
    "Parquet": {
        "extension": ".parquet",
        "filetypes": [("Parquet files", "*.parquet"), ("All files", "*.*")],
    },
}


class CancelledOperation(Exception):
    """Raised when a long-running engine operation is cancelled."""


class SourceChangedError(ValueError):
    """A source no longer matches the data approved during validation."""


def clean_cols(cols):
    return (
        pd.Index(cols)
        .astype(str)
        .str.strip()
        .str.replace("\n", " ", regex=False)
        .str.replace("\r", " ", regex=False)
        .str.replace(r"\s+", " ", regex=True)
        .str.lower()
        .tolist()
    )


def dup_cols(cols):
    duplicates, seen = [], set()
    for col in cols:
        if col in seen and col not in duplicates:
            duplicates.append(col)
        seen.add(col)
    return duplicates


def source_cache_key(source):
    if source["mode"] == "direct":
        return ("direct", source["type"], source["full_path"])
    return ("zip", source["type"], source["zip_path"], source["inner_name"])


def source_signature(source):
    path = source["full_path"] if source["mode"] == "direct" else source["zip_path"]
    info = os.stat(path)
    return (info.st_size, info.st_mtime_ns, info.st_ctime_ns)


def provenance_column(columns):
    used = set(columns)
    name = "source_file"
    if name in used:
        name = "mergeguru_source_file"
    suffix = 2
    while name in used:
        name = f"mergeguru_source_file_{suffix}"
        suffix += 1
    return name


def prepare_source_frame(source, reference_columns, sheet=None, header_row=1,
                         footer_rows=0, allow_extra=False):
    before = source_signature(source)
    approved = source.get("_validated_signature")
    if approved is not None and approved != before:
        raise SourceChangedError(f"{source['display_name']} changed after validation. Run the file check again.")
    frame = read_full(source, sheet=sheet, header_row=header_row, footer_rows=footer_rows)
    if source_signature(source) != before:
        raise SourceChangedError(f"{source['display_name']} changed while reading. Run the file check again.")
    frame.columns = clean_cols(frame.columns)
    missing = [col for col in reference_columns if col not in frame.columns]
    extra = [col for col in frame.columns if col not in reference_columns]
    duplicates = dup_cols(frame.columns)
    if missing or duplicates or (extra and not allow_extra):
        raise SourceChangedError(
            f"{source['display_name']} layout changed or is invalid: "
            f"missing={missing}, extra={extra}, duplicates={duplicates}. Run the file check again."
        )
    return frame


def _infer_source_values(frame):
    # Keep lexical IDs intact before numeric inference can strip their zeros.
    for column in frame.columns:
        series = frame[column]
        if not pd.api.types.is_object_dtype(series):
            continue
        values = series.dropna()
        if values.empty or not all(isinstance(value, str) for value in values):
            continue
        if values.str.match(r"^\s*[+-]?(?:0\d|\d{16,})").any():
            continue
        try:
            frame[column] = pd.to_numeric(series, errors="raise")
        except (ValueError, TypeError):
            if values.isin(["True", "False"]).all():
                frame[column] = series.map({"True": True, "False": False})
    return frame.infer_objects()


def discover_sources(paths):
    sources = []
    errors = []
    for raw_path in paths:
        path = os.path.abspath(raw_path)
        name = os.path.basename(path)
        suffix = os.path.splitext(path)[1].lower()
        if suffix in {".xlsx", ".csv"}:
            sources.append(
                {
                    "type": suffix[1:],
                    "mode": "direct",
                    "display_name": name,
                    "full_path": path,
                }
            )
        elif suffix == ".zip":
            try:
                with zipfile.ZipFile(path, "r") as archive:
                    for inner_name in archive.namelist():
                        inner_suffix = os.path.splitext(inner_name)[1].lower()
                        if inner_suffix in {".xlsx", ".csv"}:
                            sources.append(
                                {
                                    "type": inner_suffix[1:],
                                    "mode": "zip",
                                    "display_name": f"{name} :: {inner_name}",
                                    "zip_path": path,
                                    "inner_name": inner_name,
                                }
                            )
            except Exception as exc:
                errors.append({"path": path, "error": str(exc)})
    return sources, errors


def source_size(source):
    try:
        if source["mode"] == "direct":
            return os.path.getsize(source["full_path"])
        with zipfile.ZipFile(source["zip_path"], "r") as archive:
            return archive.getinfo(source["inner_name"]).file_size
    except Exception:
        return 0


def possible_duplicate_sources(sources):
    groups = {}
    for source in sources:
        name = os.path.basename(
            source.get("inner_name") or source.get("full_path") or source["display_name"]
        ).lower()
        groups.setdefault((name, source_size(source)), []).append(source["display_name"])
    return [names for names in groups.values() if len(names) > 1]


def read_csv_flex(source, nrows=None, header=0, skipfooter=0):
    for encoding in ("utf-8", "latin-1"):
        try:
            if hasattr(source, "seek"):
                source.seek(0)
            kwargs = {"nrows": nrows, "encoding": encoding, "header": header, "dtype": object}
            if skipfooter:
                kwargs.update({"skipfooter": skipfooter, "engine": "python"})
            return _infer_source_values(pd.read_csv(source, **kwargs))
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("utf-8", b"", 0, 1, "Unable to decode CSV")


def get_sheet_names(source):
    key = (source_cache_key(source), source_signature(source))
    if key in _SHEET_NAME_CACHE:
        return list(_SHEET_NAME_CACHE[key])
    if source["mode"] == "direct":
        workbook = load_workbook(source["full_path"], read_only=True)
    else:
        with zipfile.ZipFile(source["zip_path"], "r") as archive:
            workbook = load_workbook(
                BytesIO(archive.read(source["inner_name"])), read_only=True
            )
    try:
        names = tuple(workbook.sheetnames)
    finally:
        workbook.close()
    _SHEET_NAME_CACHE[key] = names
    return list(names)


def _read_zipped_csv(source, nrows=None, header=0, skipfooter=0):
    with zipfile.ZipFile(source["zip_path"], "r") as archive:
        with archive.open(source["inner_name"]) as stream:
            return read_csv_flex(
                stream, nrows=nrows, header=header, skipfooter=skipfooter
            )


def read_header(source, sheet=None, header_row=1, footer_rows=0):
    header_index = max(0, int(header_row) - 1)
    if source["type"] == "xlsx":
        if not sheet:
            raise ValueError(f"Sheet required for: {source['display_name']}")
        if source["mode"] == "direct":
            return pd.read_excel(
                source["full_path"],
                sheet_name=sheet,
                header=header_index,
                nrows=0,
                skipfooter=footer_rows,
            ).columns
        with zipfile.ZipFile(source["zip_path"], "r") as archive:
            return pd.read_excel(
                BytesIO(archive.read(source["inner_name"])),
                sheet_name=sheet,
                header=header_index,
                nrows=0,
                skipfooter=footer_rows,
            ).columns
    if source["mode"] == "direct":
        return read_csv_flex(
            source["full_path"],
            nrows=0,
            header=header_index,
            skipfooter=footer_rows,
        ).columns
    return _read_zipped_csv(
        source, nrows=0, header=header_index, skipfooter=footer_rows
    ).columns


def read_full(source, sheet=None, header_row=1, footer_rows=0):
    header_index = max(0, int(header_row) - 1)
    if source["type"] == "xlsx":
        if not sheet:
            raise ValueError(f"Sheet required for: {source['display_name']}")
        if source["mode"] == "direct":
            return _infer_source_values(pd.read_excel(
                source["full_path"],
                sheet_name=sheet,
                header=header_index,
                skipfooter=footer_rows,
                dtype=object,
            ))
        with zipfile.ZipFile(source["zip_path"], "r") as archive:
            return _infer_source_values(pd.read_excel(
                BytesIO(archive.read(source["inner_name"])),
                sheet_name=sheet,
                header=header_index,
                skipfooter=footer_rows,
                dtype=object,
            ))
    if source["mode"] == "direct":
        return read_csv_flex(
            source["full_path"], header=header_index, skipfooter=footer_rows
        )
    return _read_zipped_csv(
        source, header=header_index, skipfooter=footer_rows
    )


def resolve_sheet(source, reference_sheet):
    if source["type"] != "xlsx":
        return None
    names = get_sheet_names(source)
    if reference_sheet in names:
        return reference_sheet
    raise ValueError(
        f"Sheet '{reference_sheet}' not found. Available: {', '.join(names)}"
    )


def match_sheet_by_headers(
    source,
    reference_columns,
    reference_sheet=None,
    header_row=1,
    footer_rows=0,
    allow_extra=False,
):
    if source["type"] != "xlsx":
        return None, clean_cols(read_header(source, header_row=header_row, footer_rows=footer_rows)), []

    matches = []
    for sheet in get_sheet_names(source):
        columns = clean_cols(
            read_header(
                source,
                sheet=sheet,
                header_row=header_row,
                footer_rows=footer_rows,
            )
        )
        duplicates = dup_cols(columns)
        if duplicates:
            continue
        missing = [col for col in reference_columns if col not in columns]
        extra = [col for col in columns if col not in reference_columns]
        exact = not missing and not extra
        if not missing:
            matches.append((sheet, columns, exact, extra))

    if not matches:
        names = ", ".join(get_sheet_names(source))
        raise ValueError(
            "No worksheet matches the reference column layout. "
            f"Available sheets: {names}"
        )

    matches.sort(key=lambda item: (
        item[0] != reference_sheet,
        not item[2],
        get_sheet_names(source).index(item[0]),
    ))
    sheet, columns, _exact, _extra = matches[0]
    notes = [f"matched sheet: {sheet}"]
    if len(matches) > 1:
        notes.append(f"multiple matching sheets found; used {sheet}")
    return sheet, columns, notes


def classify_file_error(exc):
    message = str(exc)
    lowered = message.lower()
    if "password" in lowered or "encrypted" in lowered:
        return "password-protected", "Password-protected or encrypted file is unsupported."
    if isinstance(exc, PermissionError) or "permission denied" in lowered or "being used" in lowered:
        return "locked/permission", "File is locked or cannot be read with current permissions."
    if isinstance(exc, zipfile.BadZipFile) or "not a zip file" in lowered:
        return "corrupt", "File appears corrupt, incomplete, or invalid."
    return "read error", message or exc.__class__.__name__


def validate_sources(
    sources,
    reference_index,
    selected_sheet=None,
    header_row=1,
    footer_rows=0,
    allow_extra=False,
    auto_match_sheet=False,
    progress_callback=None,
):
    reference = sources[reference_index]
    reference_sheet = selected_sheet if reference["type"] == "xlsx" else None
    reference_columns = clean_cols(
        read_header(
            reference,
            sheet=reference_sheet,
            header_row=header_row,
            footer_rows=footer_rows,
        )
    )
    reference_duplicates = dup_cols(reference_columns)
    if reference_duplicates:
        raise ValueError(
            "Reference has duplicate columns: " + ", ".join(reference_duplicates)
        )

    def validate_one(index, source):
        missing, extra, duplicates = [], [], []
        sheet = None
        try:
            signature = source_signature(source)
            if source is reference:
                sheet = reference_sheet
                columns = clean_cols(
                    read_header(
                        source,
                        sheet=sheet,
                        header_row=header_row,
                        footer_rows=footer_rows,
                    )
                )
                sheet_notes = []
            elif source["type"] == "xlsx" and auto_match_sheet:
                sheet, columns, sheet_notes = match_sheet_by_headers(
                    source,
                    reference_columns,
                    reference_sheet=reference_sheet,
                    header_row=header_row,
                    footer_rows=footer_rows,
                    allow_extra=allow_extra,
                )
            else:
                sheet = resolve_sheet(source, reference_sheet) if source["type"] == "xlsx" else None
                columns = clean_cols(
                    read_header(
                        source,
                        sheet=sheet,
                        header_row=header_row,
                        footer_rows=footer_rows,
                    )
                )
                sheet_notes = []
            duplicates = dup_cols(columns)
            missing = [col for col in reference_columns if col not in columns]
            extra = [col for col in columns if col not in reference_columns]
            accepted_extra = allow_extra and extra and not missing and not duplicates
            passed = (not missing and not extra and not duplicates) or accepted_extra
            if signature != source_signature(source):
                raise SourceChangedError("Source changed during validation. Run the file check again.")
            if passed:
                source["_validated_signature"] = signature
            status = "warning" if accepted_extra else "pass" if passed else "fail"
            details = []
            if missing:
                details.append("missing: " + ", ".join(missing[:5]))
            if extra:
                details.append("extra: " + ", ".join(extra[:5]))
            if duplicates:
                details.append("duplicate headers: " + ", ".join(duplicates))
            if not details:
                details.append(f"{len(columns)} columns")
            details.extend(sheet_notes)
        except Exception as exc:
            category, detail = classify_file_error(exc)
            passed, status = False, "fail"
            details = [f"{category}: {detail}"]
        return {
            "index": index,
            "source": source,
            "name": source["display_name"],
            "status": status,
            "passed": passed,
            "missing": missing,
            "extra": extra,
            "duplicates": duplicates,
            "sheet": sheet,
            "detail": "; ".join(details),
        }

    results = []
    workers = min(6, max(1, len(sources)))
    with ThreadPoolExecutor(max_workers=workers) as executor:
        futures = {
            executor.submit(validate_one, index, source): index
            for index, source in enumerate(sources)
        }
        for completed, future in enumerate(as_completed(futures), start=1):
            results.append(future.result())
            if progress_callback:
                progress_callback(completed, len(sources))
    results.sort(key=lambda item: item["index"])
    passed_sources = [item["source"] for item in results if item["passed"]]
    source_sheet_map = {
        source_cache_key(item["source"]): item["sheet"]
        for item in results
        if item["passed"]
    }
    extra_columns = []
    for item in results:
        if item["passed"]:
            for col in item["extra"]:
                if col not in extra_columns:
                    extra_columns.append(col)
    summary = {
        "pass": sum(item["status"] == "pass" for item in results),
        "warn": sum(item["status"] == "warning" for item in results),
        "fail": sum(item["status"] == "fail" for item in results),
        "missing": sum(bool(item["missing"]) for item in results),
        "extra": sum(bool(item["extra"]) for item in results),
        "duplicate": sum(bool(item["duplicates"]) for item in results),
    }
    return {
        "reference_columns": reference_columns,
        "passed_sources": passed_sources,
        "source_sheet_map": source_sheet_map,
        "extra_columns": extra_columns,
        "results": results,
        "summary": summary,
    }


def choose_combine_workers(sources, total_bytes):
    if len(sources) <= 1:
        return 1
    cap = min(4, max(1, (os.cpu_count() or 2) // 2), len(sources))
    if any(source["type"] == "xlsx" for source in sources):
        cap = min(cap, 2)
    if total_bytes >= 1024 ** 3:
        cap = min(cap, 2)
    return max(1, cap)


def combine_sources(
    sources,
    reference_source,
    reference_columns,
    selected_sheet=None,
    source_sheet_map=None,
    header_row=1,
    footer_rows=0,
    allow_extra=False,
    extra_columns=None,
    fast_mode=False,
    optimize_memory=False,
    progress_callback=None,
    cancel_callback=None,
):
    ordered_extras = list(extra_columns or [])
    total_bytes = sum(source_size(source) for source in sources)
    source_sheet_map = source_sheet_map or {}

    def load_one(index, source):
        if cancel_callback and cancel_callback():
            raise CancelledOperation("Combine cancelled.")
        if source["type"] != "xlsx":
            sheet = None
        else:
            sheet = source_sheet_map.get(source_cache_key(source))
            if sheet is None:
                sheet = selected_sheet if source is reference_source else resolve_sheet(source, selected_sheet)
        frame = prepare_source_frame(
            source,
            reference_columns,
            sheet=sheet,
            header_row=header_row,
            footer_rows=footer_rows,
            allow_extra=allow_extra,
        )
        return index, source, frame, source_size(source)

    results = []
    bytes_done = 0
    if fast_mode and len(sources) > 1:
        workers = choose_combine_workers(sources, total_bytes)
        with ThreadPoolExecutor(max_workers=workers) as executor:
            futures = {
                executor.submit(load_one, index, source): source
                for index, source in enumerate(sources)
            }
            for completed, future in enumerate(as_completed(futures), start=1):
                if cancel_callback and cancel_callback():
                    raise CancelledOperation("Combine cancelled.")
                result = future.result()
                results.append(result)
                bytes_done += result[3]
                if progress_callback:
                    progress_callback(completed, len(sources), bytes_done, total_bytes)
    else:
        for index, source in enumerate(sources):
            result = load_one(index, source)
            results.append(result)
            bytes_done += result[3]
            if progress_callback:
                progress_callback(index + 1, len(sources), bytes_done, total_bytes)

    results.sort(key=lambda item: item[0])
    if allow_extra:
        for _, _, frame, _ in results:
            for col in frame.columns:
                if col not in reference_columns and col not in ordered_extras:
                    ordered_extras.append(col)
    final_columns = list(reference_columns)
    if allow_extra:
        final_columns.extend(col for col in ordered_extras if col not in final_columns)
    provenance = provenance_column(final_columns)
    for _, source, frame, _ in results:
        frame[provenance] = source["display_name"]
    final_columns.append(provenance)
    frames = [frame.reindex(columns=final_columns) for _, _, frame, _ in results]
    if not frames:
        return pd.DataFrame(columns=final_columns), ordered_extras, None
    rolling = [
        pd.concat(frames[start : start + 10], ignore_index=True)
        for start in range(0, len(frames), 10)
    ]
    final = pd.concat(rolling, ignore_index=True)
    if cancel_callback and cancel_callback():
        raise CancelledOperation("Combine cancelled.")
    memory_stats = None
    if optimize_memory:
        final, memory_stats = optimize_dataframe_memory(final)
    return final, ordered_extras, memory_stats


def build_filter_mask(df, filter_config, string_series=None, date_series=None):
    column = filter_config["column"]
    mode = filter_config["mode"]
    if mode == "category":
        series = string_series if string_series is not None else df[column].fillna("").astype(str)
        matched = series.isin(filter_config["values"])
        return matched if filter_config["action"] == "include" else ~matched
    if mode == "text":
        series = string_series if string_series is not None else df[column].fillna("").astype(str)
        operator = filter_config["operator"]
        value = filter_config.get("value", "")
        if operator == "equals":
            return series == value
        if operator == "not_equals":
            return series != value
        if operator == "contains":
            return series.str.contains(value, case=False, na=False)
        if operator == "starts_with":
            return series.str.startswith(value, na=False)
        if operator == "paste_include":
            return series.isin(filter_config["values"])
        if operator == "paste_exclude":
            return ~series.isin(filter_config["values"])
    if mode == "date":
        dates = date_series if date_series is not None else pd.to_datetime(df[column], errors="coerce")
        first = pd.to_datetime(filter_config["date1"], errors="raise")
        operator = filter_config["operator"]
        if operator == "equals":
            return dates.dt.date == first.date()
        if operator == "before":
            return dates < first
        if operator == "after":
            return dates > first
        second = pd.to_datetime(filter_config["date2"], errors="raise")
        return (dates >= first) & (dates <= second)
    return pd.Series(True, index=df.index)


def apply_filters(df, filters):
    if df is None or not filters:
        return df
    mask = pd.Series(True, index=df.index)
    string_cache = {}
    date_cache = {}
    for config in filters:
        column = config["column"]
        string_series = None
        date_series = None
        if config["mode"] in {"category", "text"}:
            string_series = string_cache.setdefault(
                column, df[column].fillna("").astype(str)
            )
        elif config["mode"] == "date":
            date_series = date_cache.setdefault(
                column, pd.to_datetime(df[column], errors="coerce")
            )
        mask &= build_filter_mask(df, config, string_series, date_series)
    return df.loc[mask]


def filter_summary(config):
    column, mode = config["column"], config["mode"]
    if mode == "category":
        values = ", ".join(config["values"][:3])
        if len(config["values"]) > 3:
            values += "..."
        return f"{column} > {config['action'].upper()} [{values}]"
    if mode == "text":
        operator = config["operator"].replace("_", " ")
        value = config.get("value") or ", ".join(config.get("values", [])[:3])
        return f"{column} > {operator}: {value}"
    if mode == "date":
        if config["operator"] == "between":
            return f"{column} > between {config['date1']} - {config['date2']}"
        return f"{column} > {config['operator']} {config['date1']}"
    return str(config)


def detect_export_format(path, fallback="Excel"):
    extension = os.path.splitext(path)[1].lower()
    if extension == ".csv" and fallback == "CSV-Unique text":
        return fallback
    for name, metadata in EXPORT_FORMATS.items():
        if name != "CSV-Unique text" and extension == metadata["extension"]:
            return name
    return fallback


def ensure_export_extension(path, export_format):
    base, extension = os.path.splitext(path)
    wanted = EXPORT_FORMATS[export_format]["extension"]
    if extension.lower() == wanted:
        return path
    return (base if extension else path) + wanted


def save_dataframe(
    df,
    path,
    export_format,
    unique_column=None,
    unique_columns=None,
    validation_report_df=None,
    progress_callback=None,
    cancel_callback=None,
):
    if export_format not in EXPORT_FORMATS:
        raise ValueError(f"Unsupported export format: {export_format}")
    if export_format == "Excel":
        for frame in (df, validation_report_df):
            if frame is None:
                continue
            if len(frame) + 1 > EXCEL_HARD_ROW_LIMIT:
                raise ValueError("Excel row limit exceeded, including the header. Choose CSV or Parquet.")
            if len(frame.columns) > EXCEL_HARD_COLUMN_LIMIT:
                raise ValueError("Excel column limit exceeded. Choose CSV or Parquet.")
    if cancel_callback and cancel_callback():
        raise CancelledOperation("Export cancelled.")

    # Commit only a completed export, preserving any previous destination.
    destination = os.path.abspath(path)
    descriptor, temporary = tempfile.mkstemp(
        prefix=".mergeguru-", suffix=EXPORT_FORMATS[export_format]["extension"],
        dir=os.path.dirname(destination),
    )
    os.close(descriptor)
    try:
        _write_dataframe(df, temporary, export_format, unique_column, unique_columns,
                         validation_report_df, progress_callback, cancel_callback)
        if cancel_callback and cancel_callback():
            raise CancelledOperation("Export cancelled.")
        os.replace(temporary, destination)
    finally:
        if os.path.exists(temporary):
            os.remove(temporary)
    if progress_callback:
        progress_callback(100, f"Completed: {len(df):,} rows")


def _excel_cell_value(value):
    """Convert pandas/numpy scalar values into safe XlsxWriter values."""
    if value is None:
        return None
    try:
        missing = pd.isna(value)
        if bool(missing):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, pd.Timestamp):
        if value.tzinfo is not None:
            value = value.tz_convert("UTC").tz_localize(None)
        return value.to_pydatetime()
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.astimezone(timezone.utc).replace(tzinfo=None)
    if value.__class__.__module__.startswith("numpy") and hasattr(value, "item"):
        value = value.item()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    if isinstance(value, (dict, list, set, tuple)):
        return str(value)
    return value


def _write_excel_streaming(df, path, validation_report_df=None,
                           progress_callback=None, cancel_callback=None):
    """Write rows in order without retaining an in-memory workbook cell tree."""
    sheets = [("Data", df)]
    if validation_report_df is not None and not validation_report_df.empty:
        sheets.append(("Validation Report", validation_report_df))
    total_rows = sum(len(frame) for _, frame in sheets)
    written = 0
    options = {
        "constant_memory": True,
        "default_date_format": "yyyy-mm-dd hh:mm:ss",
        "strings_to_formulas": False,
        "strings_to_urls": False,
        "use_zip64": True,
    }
    with xlsxwriter.Workbook(path, options) as workbook:
        header_format = workbook.add_format({"bold": True})
        for sheet_name, frame in sheets:
            worksheet = workbook.add_worksheet(sheet_name)
            worksheet.write_row(0, 0, [str(column) for column in frame.columns], header_format)
            worksheet.freeze_panes(1, 0)
            sheet_rows = len(frame)
            for row_number, row in enumerate(
                frame.itertuples(index=False, name=None), start=1
            ):
                if row_number == 1 or row_number % 1_000 == 0:
                    if cancel_callback and cancel_callback():
                        raise CancelledOperation("Export cancelled.")
                worksheet.write_row(
                    row_number,
                    0,
                    [_excel_cell_value(value) for value in row],
                )
                written += 1
                if written % 10_000 == 0 or row_number == sheet_rows:
                    if progress_callback:
                        progress_callback(
                            min(99, int(written / max(1, total_rows) * 99)),
                            f"Writing {sheet_name}: {written:,} of {total_rows:,} rows",
                        )
                    if cancel_callback and cancel_callback():
                        raise CancelledOperation("Export cancelled.")


def _write_dataframe(df, path, export_format, unique_column, unique_columns,
                     validation_report_df, progress_callback, cancel_callback):
    if export_format in {"CSV", "CSV-Unique text"}:
        working = df
        if export_format == "CSV-Unique text":
            working = df.copy()
            protected_columns = []
            if unique_columns:
                protected_columns.extend(unique_columns)
            elif unique_column:
                protected_columns.append(unique_column)
            for column in protected_columns:
                if column in working.columns:
                    values = working[column]
                    working[column] = ("'" + values.astype(str) + "'").where(values.notna(), "")
        total = len(working)
        chunk_size = max(10_000, total // 20) if total else 10_000
        with open(path, "w", newline="", encoding="utf-8") as stream:
            working.head(0).to_csv(stream, index=False)
            for start in range(0, total, chunk_size):
                if cancel_callback and cancel_callback():
                    raise CancelledOperation("Export cancelled.")
                working.iloc[start : start + chunk_size].to_csv(
                    stream, index=False, header=False
                )
                if progress_callback:
                    progress_callback(
                        min(99, int((start + chunk_size) / max(1, total) * 100)),
                        f"Writing {min(start + chunk_size, total):,} of {total:,} rows",
                    )
        return
    if cancel_callback and cancel_callback():
        raise CancelledOperation("Export cancelled.")
    if export_format == "Excel":
        _write_excel_streaming(
            df,
            path,
            validation_report_df=validation_report_df,
            progress_callback=progress_callback,
            cancel_callback=cancel_callback,
        )
    elif export_format == "Parquet":
        df.to_parquet(path, index=False)
    else:
        raise ValueError(f"Unsupported export format: {export_format}")


def get_available_ram_gb():
    try:
        return psutil.virtual_memory().available / (1024 ** 3)
    except Exception:
        return None


def get_process_ram_usage_mb():
    try:
        return psutil.Process(os.getpid()).memory_info().rss / (1024 ** 2)
    except Exception:
        return None


def get_dataframe_memory_mb(df):
    try:
        return df.memory_usage(deep=True).sum() / (1024 ** 2)
    except Exception:
        return None


def optimize_dataframe_memory(df):
    before = get_dataframe_memory_mb(df)
    stats = {"int_cols": 0, "float_cols": 0, "category_cols": 0}
    for col in df.columns:
        series = df[col]
        try:
            if pd.api.types.is_integer_dtype(series):
                df[col] = pd.to_numeric(series, downcast="integer")
                stats["int_cols"] += 1
            elif pd.api.types.is_float_dtype(series):
                candidate = pd.to_numeric(series, downcast="float")
                if candidate.dtype != series.dtype and candidate.astype(series.dtype).equals(series):
                    df[col] = candidate
                    stats["float_cols"] += 1
            elif pd.api.types.is_object_dtype(series):
                non_null = series.dropna()
                if not non_null.empty and non_null.nunique() <= min(1000, max(1, len(series) // 2)):
                    df[col] = series.astype("category")
                    stats["category_cols"] += 1
        except Exception:
            continue
    after = get_dataframe_memory_mb(df)
    stats.update(
        {
            "before_mb": before,
            "after_mb": after,
            "saved_mb": max(0.0, before - after)
            if before is not None and after is not None
            else None,
        }
    )
    return df, stats


def check_memory_before_combine(total_file_size_bytes, _num_files):
    available = get_available_ram_gb()
    if available is None:
        return True, "", None, None
    estimated = total_file_size_bytes * 1.3 / (1024 ** 3)
    safe_threshold = estimated * 2 + 1.5
    safe = available > safe_threshold
    warning = ""
    if not safe:
        warning = (
            f"Available RAM: {available:.1f} GB\n"
            f"Estimated requirement: {estimated:.1f} GB\n\n"
            "Close heavy applications, process fewer files, or use sequential mode."
        )
    return safe, warning, available, estimated
